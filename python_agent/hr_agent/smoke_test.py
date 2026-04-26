"""Smoke-тест: проверяет что агент работает на всех уровнях.

Запуск:
    PYTHONIOENCODING=utf-8 python -m hr_agent.smoke_test

Что проверяется:
  [1/4] Конфиг подхватился (USE_MOCK или ключ есть)
  [2/4] Прямой вызов 3 функций агента (start/next/evaluate)
  [3/4] Парсер резюме (DOCX)
  [4/4] HTTP-эндпоинты (если flask установлен) — поднимает сервер на 5050,
        дёргает Go-style /api/questions и /api/score, проверяет формат, гасит сервер
"""
from __future__ import annotations

import json
import sys
import threading
import time
import urllib.error
import urllib.request


def step(n: int, total: int, title: str) -> None:
    print(f"\n[{n}/{total}] {title}")
    print("-" * 60)


def ok(msg: str) -> None:
    print(f"  [OK] {msg}")


def fail(msg: str) -> None:
    print(f"  [FAIL] {msg}")
    sys.exit(1)


# ---------- 1. Конфиг ----------

def check_config() -> None:
    step(1, 4, "Конфигурация")
    from . import config
    print(f"  USE_MOCK = {config.USE_MOCK}")
    print(f"  GIGACHAT_AUTH_KEY length = {len(config.GIGACHAT_AUTH_KEY)}")
    print(f"  Model = {config.GIGACHAT_MODEL}")
    if config.USE_MOCK:
        ok("Mock-режим — реальный API не дёргается")
    else:
        if not config.GIGACHAT_AUTH_KEY:
            fail("USE_MOCK=false, но GIGACHAT_AUTH_KEY пустой!")
        ok(f"Реальный режим — будет дёргаться GigaChat ({config.GIGACHAT_MODEL})")


# ---------- 2. Прямые вызовы функций агента ----------

def check_agent_functions() -> None:
    step(2, 4, "Функции агента (start/next/evaluate)")
    from . import start_interview, next_question, evaluate_candidate

    profile = {
        "name": "Smoke Test",
        "position": "Python-разработчик",
        "grade": "Middle",
        "about": "Бэкендер 3 года",
        "education": "МГТУ",
        "resume": "Python, FastAPI, PostgreSQL, Docker, 3 года в продакшене",
    }

    print("  Запускаю start_interview...")
    start = start_interview(profile)
    assert "greeting" in start and "first_question" in start
    ok(f"start_interview: первый вопрос — '{start['first_question'][:60]}...'")

    dialogue = f"HR: {start['first_question']}\nКандидат: Делал сервис на FastAPI\n"
    ctx = start["session_context"]

    print("  Запускаю next_question...")
    step2 = next_question(profile, dialogue, ctx)
    assert "is_finished" in step2
    ok(f"next_question: {'завершено' if step2['is_finished'] else step2['question'][:60]}...")

    if not step2["is_finished"]:
        dialogue += f"HR: {step2['question']}\nКандидат: Английский B2\n"

    print("  Запускаю evaluate_candidate...")
    result = evaluate_candidate(profile, dialogue)
    assert "score" in result
    assert "scoreNormalized" in result
    assert "status" in result
    assert "extractedFacts" in result
    assert 0 <= result["scoreNormalized"] <= 1
    assert result["status"] in {"Recommended", "Maybe", "Not recommended"}
    ok(f"evaluate_candidate: score={result['score']}, status={result['status']}")
    ok(f"  extractedFacts.englishLevel = {result['extractedFacts'].get('englishLevel')}")


# ---------- 3. Парсер резюме ----------

def check_parser() -> None:
    step(3, 4, "Парсер резюме (DOCX)")
    try:
        from docx import Document
    except ImportError:
        fail("python-docx не установлен — pip install python-docx")
    from pathlib import Path
    from . import parse_resume_file

    doc = Document()
    doc.add_heading("Резюме", level=1)
    doc.add_paragraph("Иван Петров, Python-разработчик")
    doc.add_paragraph("Опыт: FastAPI, PostgreSQL, 3 года")
    tmp = Path("__smoke_resume.docx")
    doc.save(tmp)
    try:
        text = parse_resume_file(tmp)
        assert "Иван Петров" in text
        ok(f"Извлечено {len(text)} символов из DOCX, кириллица читается")
    finally:
        tmp.unlink(missing_ok=True)


# ---------- 4. HTTP-эндпоинты ----------

def check_http_endpoints() -> None:
    step(4, 4, "HTTP-эндпоинты (Go-совместимый формат)")
    try:
        from .api_example import app
    except ImportError:
        print("  [SKIP] Flask не установлен — pip install flask")
        return

    port = 5050  # отдельный порт чтобы не конфликтовать с реальным сервером
    server_thread = threading.Thread(
        target=lambda: app.run(host="127.0.0.1", port=port, debug=False, use_reloader=False),
        daemon=True,
    )
    server_thread.start()

    base = f"http://127.0.0.1:{port}"
    for _ in range(20):
        try:
            urllib.request.urlopen(f"{base}/api/health", timeout=1).read()
            break
        except urllib.error.URLError:
            time.sleep(0.3)
    else:
        fail("Сервер не поднялся за 6 секунд")

    # Health
    health = json.loads(urllib.request.urlopen(f"{base}/api/health").read())
    assert "ok" in health and health["ok"]
    ok(f"/api/health → ok=true, use_mock={health.get('use_mock')}")

    # Go-style /api/questions
    req = urllib.request.Request(
        f"{base}/api/questions",
        data=json.dumps({"position": "backend_go", "grade": "Middle"}).encode("utf-8"),
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    questions_resp = json.loads(urllib.request.urlopen(req).read())
    assert "questions" in questions_resp
    questions = questions_resp["questions"]
    assert len(questions) > 0
    assert all("id" in q and "text" in q for q in questions)
    ok(f"POST /api/questions → {len(questions)} вопросов в формате [{{id,text}}]")

    # Go-style /api/score
    payload = {
        "candidate": {
            "candidate_id": "smoke_test",
            "position": "backend_go",
            "grade": "Middle",
            "about": "Go developer 3 years",
            "education": "MGTU",
            "resume_url": "",
        },
        "answers": [
            {"QuestionID": 1, "QuestionText": "Опыт", "AnswerText": "Делал сервис на Go, 5к RPS"},
            {"QuestionID": 2, "QuestionText": "Английский", "AnswerText": "B2"},
            {"QuestionID": 3, "QuestionText": "Зарплата", "AnswerText": "250к"},
        ],
    }
    req = urllib.request.Request(
        f"{base}/api/score",
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    score_resp = json.loads(urllib.request.urlopen(req).read())

    # Проверка Go-контракта
    required = ["score", "score_percent", "verdict", "recommendation",
                "traits", "rejection_timer_days", "rejection_message_template", "status"]
    for field in required:
        if field not in score_resp:
            fail(f"В ответе /api/score нет обязательного поля '{field}'")
    assert 0 <= score_resp["score"] <= 1, "score должен быть 0..1"
    assert 0 <= score_resp["score_percent"] <= 100
    assert score_resp["verdict"] in {"recommend_to_interview", "soft_reject_timer"}
    assert score_resp["status"] in {"approved", "pending_rejection"}
    assert "positive" in score_resp["traits"] and "negative" in score_resp["traits"]

    ok(f"POST /api/score → score={score_resp['score']}, "
       f"verdict={score_resp['verdict']}, status={score_resp['status']}")
    ok(f"  traits.positive: {len(score_resp['traits']['positive'])} штук")
    ok(f"  traits.negative: {len(score_resp['traits']['negative'])} штук")
    if "extra" in score_resp:
        ok("  extra-поля для фронта: есть")


# ---------- main ----------

def main() -> None:
    print("=" * 60)
    print(" SMOKE TEST — HR-агент на GigaChat")
    print("=" * 60)
    try:
        check_config()
        check_agent_functions()
        check_parser()
        check_http_endpoints()
    except AssertionError as exc:
        fail(f"Assertion failed: {exc}")

    print("\n" + "=" * 60)
    print(" [SUCCESS] Все проверки пройдены")
    print("=" * 60)


if __name__ == "__main__":
    main()
